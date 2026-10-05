"""Inspect structures and mechanically fill Agent supplied mappings."""
import copy
import hashlib
import json
import os
import re
import shutil
import zipfile
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter, column_index_from_string
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from doc_extract import extract_document
import lark_io
import source_trace

def detect_format(path):
    kind = extract_document(path)["format"]
    if kind == 'xlsx':
        with zipfile.ZipFile(path) as archive:
            if 'xl/vbaProject.bin' in archive.namelist():
                return 'xlsm'
    return kind


def protect_output(out, inputs=(), overwrite=False):
    target = Path(out).expanduser().resolve()
    for src in inputs:
        source = Path(src).expanduser().resolve()
        if target == source or target.is_relative_to(source.parent):
            raise ValueError('Output must be outside every input directory')
    if target.exists() and not overwrite:
        raise ValueError('Output already exists: ' + str(target))
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def inspect_template(path):
    fmt = detect_format(path)
    if fmt in ('xlsx', 'xlsm'):
        wb = load_workbook(path, keep_vba=fmt == 'xlsm')
        result = {'format': fmt, 'sheets': []}
        for ws in wb:
            result['sheets'].append({'name': ws.title, 'rows': ws.max_row, 'columns': ws.max_column,
                'preview': [{'coordinate': c.coordinate, 'value': c.value} for row in ws.iter_rows(max_row=min(30, ws.max_row)) for c in row if c.value is not None],
                'merges': [str(r) for r in ws.merged_cells.ranges],
                'column_widths': {k:v.width for k,v in ws.column_dimensions.items()},
                'row_heights': {str(k):v.height for k,v in ws.row_dimensions.items()},
                'freeze_panes': str(ws.freeze_panes) if ws.freeze_panes else None,
                'validations': [str(v) for v in ws.data_validations.dataValidation],
                'formulas': [{'coordinate': c.coordinate, 'value': c.value} for row in ws for c in row if c.data_type == 'f']})
        wb.close()
        return result
    return extract_document(path)


def text_value(value, mapping, numbered=False):
    if isinstance(value, list):
        value = mapping.get('list_join', '\n').join((str(i+1) + '. ' if numbered else '') + str(v) for i,v in enumerate(value))
    elif isinstance(value, dict):
        value = json.dumps(value, ensure_ascii=False)
    value = '' if value is None else str(value)
    if len(value) > 32767:
        raise ValueError('Excel cell text exceeds 32767 characters')
    return value


def field_text(case, field, mapping, numbered=False):
    value = case[field]
    if field == '参考来源':
        # Structured or legacy sources render as readable text, never hidden JSON.
        value = source_trace.cell_text(value)
    return text_value(value, mapping, numbered)


def check_fields(cases, columns):
    for c in cases:
        for field in columns.values():
            if field not in c:
                raise ValueError('Mapping references absent field: ' + field)


def set_text(cell, value):
    cell.value = value
    # User text that begins with = must stay text, never become a formula.
    cell.data_type = 's'


def fill_xlsx(template, cases_doc, mapping, out_path, overwrite=False):
    out = protect_output(out_path, [template] if template else [], overwrite)
    if template and detect_format(template) not in ('xlsx', 'xlsm'):
        raise ValueError('Non-XLSX template: inspect first and select a blank workbook mapping')
    wb = load_workbook(template, keep_vba=detect_format(template) == 'xlsm') if template else Workbook()
    creating = mapping.get('create_sheet', False)
    if type(creating) is not bool:
        raise ValueError('create_sheet must be Boolean')
    if creating and mapping.get('output_sheet'):
        raise ValueError('Choose create_sheet or output_sheet, not both')
    if (not isinstance(mapping['sheet'], str) or not mapping['sheet'] or len(mapping['sheet']) > 31
            or re.search(r'[\\/*?:\[\]]', mapping['sheet'])):
        raise ValueError('Invalid sheet name')
    if not template:
        wb.active.title = mapping['sheet']
    elif creating:
        if mapping['sheet'].lower() in {name.lower() for name in wb.sheetnames}:
            raise ValueError('New sheet already exists: ' + mapping['sheet'])
        wb.create_sheet(mapping['sheet'])
    ws = wb[mapping['sheet']]
    blank_sheet = not template or creating
    output_sheet = mapping.get('output_sheet')
    if output_sheet is not None:
        if (not isinstance(output_sheet, str) or not output_sheet or len(output_sheet) > 31
                or re.search(r'[\\/*?:\[\]]', output_sheet)):
            raise ValueError('Invalid output_sheet name')
        if output_sheet.lower() != ws.title.lower() and output_sheet.lower() in {name.lower() for name in wb.sheetnames}:
            raise ValueError('output_sheet already exists: ' + output_sheet)
        if output_sheet != ws.title and output_sheet.lower() == ws.title.lower():
            raise ValueError('Case-only sheet rename is not supported; use a distinct name')
        ws.title = output_sheet
    start = mapping['start_row']
    style_row = mapping.get('style_row', start)
    if type(start) is not int or start < 1 or type(style_row) is not int or style_row < 1:
        raise ValueError('Rows must be positive integers')
    columns = mapping['columns']
    cases = cases_doc['cases']
    check_fields(cases, columns)
    sources_sheet = sources_sheet_name(mapping, cases_doc, wb)
    mode = mapping.get('row_mode', 'per_case')
    if mode not in ('per_case', 'per_step'):
        raise ValueError('Unknown row_mode')
    counts = []
    for case in cases:
        steps, expected = case.get('测试步骤', []), case.get('预期结果', [])
        if mode == 'per_step':
            if not isinstance(steps, list) or not isinstance(expected, list) or len(steps) != len(expected):
                raise ValueError('per_step needs equal-length step and expected arrays')
            counts.append(max(1, len(steps)))
        else:
            counts.append(1)
    styles = {col: copy.copy(ws[f'{col}{style_row}']._style) for col in columns}
    height = mapping.get('row_height', ws.row_dimensions[style_row].height)
    if height is not None and (type(height) not in (int, float) or not 0 < height < float('inf')):
        raise ValueError('row_height must be a positive finite number')
    alignment_values = mapping.get('alignment', {})
    if not isinstance(alignment_values, dict):
        raise ValueError('alignment must be an object of openpyxl Alignment properties')
    # Validate before changing the workbook. Blank workbooks have readable defaults.
    Alignment(**alignment_values)
    widths = mapping.get('column_widths', {})
    for col, width in widths.items():
        column_index_from_string(col)
        if type(width) not in (int, float) or not 0 < width < float('inf'):
            raise ValueError('column_widths values must be positive finite numbers')
    clear = mapping.get('clear_example_rows', [])
    if clear:
        if not (isinstance(clear, list) and len(clear) == 2 and all(type(v) is int for v in clear) and start <= clear[0] <= clear[1]):
            raise ValueError('clear_example_rows must be [first,last] at or below start_row')
    last = start + sum(counts) - 1
    column_numbers = {column_index_from_string(col) for col in columns}
    conflicts = []
    for merged in ws.merged_cells.ranges:
        intersects = (merged.min_row <= last and merged.max_row >= start
                      and any(merged.min_col <= col <= merged.max_col for col in column_numbers))
        explicitly_cleared = clear and clear[0] <= merged.min_row and merged.max_row <= clear[1]
        if intersects and not explicitly_cleared:
            conflicts.append('merged ' + str(merged))
    for row in range(start, last + 1):
        if clear and clear[0] <= row <= clear[1]:
            continue
        for col in columns:
            cell = ws[f'{col}{row}']
            if cell.value is not None:
                conflicts.append('non-empty ' + cell.coordinate)
    if conflicts:
        wb.close()
        raise ValueError('Template fill conflicts outside clear_example_rows: ' + ', '.join(conflicts[:10]))
    if clear:
        for merged in list(ws.merged_cells.ranges):
            if merged.min_row <= clear[1] and merged.max_row >= clear[0]:
                if merged.min_row < clear[0] or merged.max_row > clear[1]:
                    raise ValueError('Clear range intersects a partial merged region')
                ws.unmerge_cells(str(merged))
        for row in ws.iter_rows(min_row=clear[0], max_row=clear[1]):
            for c in row:
                c.value = None
    row_num = start
    merges = []
    for case, count in zip(cases, counts):
        steps, expected = case.get('测试步骤', []), case.get('预期结果', [])
        for offset in range(count):
            for col, field in columns.items():
                if mode == 'per_step' and field in ('测试步骤', '预期结果'):
                    value = case[field]
                    value = value[offset] if offset < len(value) else ''
                    if mapping.get('number_steps') and value:
                        value = str(offset+1) + '. ' + str(value)
                else:
                    value = field_text(case, field, mapping, mapping.get('number_steps', False) and field in ('测试步骤', '预期结果'))
                cell = ws[f'{col}{row_num+offset}']
                cell._style = copy.copy(styles[col])
                alignment = copy.copy(cell.alignment)
                if blank_sheet:
                    alignment.wrap_text = True
                    alignment.vertical = 'top'
                for key, val in alignment_values.items():
                    setattr(alignment, key, val)
                cell.alignment = alignment
                set_text(cell, text_value(value, mapping))
            ws.row_dimensions[row_num+offset].height = height
        if count > 1:
            for col in mapping.get('merge_case_columns', []):
                if col not in columns or columns[col] in ('测试步骤', '预期结果'):
                    raise ValueError('Only mapped case-level columns may merge')
                region = f'{col}{row_num}:{col}{row_num+count-1}'
                ws.merge_cells(region)
                merges.append(region)
        row_num += count
    fixed = dict(cases_doc.get('document', {}).get('header_values', {}))
    fixed.update(mapping.get('cells', {}))
    for coordinate, value in fixed.items():
        if any(ws[coordinate].coordinate in merged and ws[coordinate].coordinate != ws.cell(merged.min_row, merged.min_col).coordinate
               for merged in ws.merged_cells.ranges):
            raise ValueError('Fixed cell is not a merged anchor: ' + coordinate)
        set_text(ws[coordinate], text_value(value, mapping))
        alignment = copy.copy(ws[coordinate].alignment)
        if blank_sheet:
            alignment.wrap_text = True
            alignment.vertical = 'top'
        for key, val in alignment_values.items():
            setattr(alignment, key, val)
        ws[coordinate].alignment = alignment
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    if sources_sheet:
        write_sources_sheet(wb.create_sheet(sources_sheet), cases_doc)
    wb.save(out)
    wb.close()
    return {'rows_written': row_num-start, 'merges': merges, 'out': str(out), 'sha256': hashlib.sha256(out.read_bytes()).hexdigest(),
            'sources': dict(source_trace.summary(cases_doc), location=sources_sheet or source_location(mapping, cases_doc, 'sources_sheet'))}


def source_location(mapping, cases_doc, key):
    if not source_trace.has_content(cases_doc):
        return 'none_recorded'
    # Only an explicit mapping choice omits the readable source table.
    return 'omitted_by_mapping' if mapping.get(key, True) is False else None


def sources_sheet_name(mapping, cases_doc, wb):
    name = mapping.get('sources_sheet', '来源与参考')
    if name is False or not source_trace.has_content(cases_doc):
        return None
    if (not isinstance(name, str) or not name or len(name) > 31 or re.search(r'[\\/*?:\[\]]', name)):
        raise ValueError('sources_sheet must be a valid sheet name or false')
    if name.lower() in {sheet.lower() for sheet in wb.sheetnames} or name.lower() == str(mapping.get('output_sheet', '')).lower():
        raise ValueError('sources_sheet already exists: ' + name + '; set another sources_sheet name')
    return name


def write_sources_sheet(ws, cases_doc):
    ws.append(source_trace.sheet_header())
    link_columns = {i + 1 for i, key in enumerate(source_trace.SHEET_COLUMNS) if key in source_trace.URL_KEYS}
    for row in source_trace.sheet_rows(cases_doc):
        ws.append([None] * len(row))
        for col, value in enumerate(row, 1):
            cell = ws.cell(ws.max_row, col)
            set_text(cell, text_value(value, {}))
            if col in link_columns and source_trace._is_url(value):
                cell.hyperlink = value
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    for col, key in enumerate(source_trace.SHEET_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col)].width = source_trace.COLUMN_WIDTHS[key]
    ws.freeze_panes = 'A2'


def fill_docx(template, cases_doc, mapping, out_path, overwrite=False):
    out = protect_output(out_path, [template], overwrite)
    doc = Document(template)
    cases = cases_doc['cases']
    check_fields(cases, mapping.get('columns', {}))
    include_sources = mapping.get('sources_section', True) is not False and source_trace.has_content(cases_doc)
    if 'table' in mapping:
        table = doc.tables[mapping['table']]
        style = copy.deepcopy(table.rows[mapping.get('style_row', len(table.rows)-1)]._tr)
        for case in cases:
            row_xml = copy.deepcopy(style)
            table._tbl.append(row_xml)
            row = table.rows[-1]
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.text = ''
            for column, field in mapping['columns'].items():
                cell = row.cells[int(column)]
                value = field_text(case, field, mapping, mapping.get('number_steps', False) and field in ('测试步骤', '预期结果'))
                paragraph = cell.paragraphs[0]
                if paragraph.runs:
                    paragraph.runs[0].text = value
                else:
                    paragraph.add_run(value)
    else:
        indices = mapping['paragraphs']
        if (not isinstance(indices, list) or not indices or any(type(i) is not int or i < 0 or i >= len(doc.paragraphs) for i in indices)
                or indices != list(range(indices[0], indices[0] + len(indices)))):
            raise ValueError('paragraphs must select one contiguous block in document order')
        selected = [doc.paragraphs[i]._p for i in indices]
        parent = selected[0].getparent()
        positions = [parent.index(element) for element in selected]
        if positions != list(range(positions[0], positions[0] + len(positions))):
            raise ValueError('Selected paragraph block cannot cross tables or other document elements')
        originals = [copy.deepcopy(element) for element in selected]
        fields = set()
        for original in originals:
            nodes = [node for node in original.iter() if node.tag.endswith('}t')]
            paragraph_fields = re.findall(r'\{\{([^{}]+)\}\}', ''.join(node.text or '' for node in nodes))
            node_fields = [field for node in nodes for field in re.findall(r'\{\{([^{}]+)\}\}', node.text or '')]
            if paragraph_fields != node_fields:
                raise ValueError('Placeholder spans text runs; use one run per placeholder')
            fields.update(paragraph_fields)
        check_fields(cases, {field: field for field in fields})
        position = parent.index(selected[0])
        for element in selected:
            parent.remove(element)
        for case in cases:
            for original in originals:
                element = copy.deepcopy(original)
                for text in element.iter():
                    if text.tag.endswith('}t') and text.text:
                        text.text = re.sub(r'\{\{([^{}]+)\}\}', lambda match: field_text(case, match.group(1), mapping), text.text)
                parent.insert(position, element)
                position += 1
    if include_sources:
        append_docx_sources(doc, cases_doc)
    doc.save(out)
    return {'out': str(out), 'cases': len(cases), 'sha256': hashlib.sha256(out.read_bytes()).hexdigest(),
            'sources': dict(source_trace.summary(cases_doc), location='appended_section' if include_sources
                            else source_location(mapping, cases_doc, 'sources_section'))}


def append_docx_sources(doc, cases_doc):
    # Plain paragraphs avoid depending on template heading/table styles.
    doc.add_paragraph().add_run('来源与参考').bold = True
    doc.add_paragraph('以下来源由编写 Agent 记录并复核；标“未提供”的字段没有可靠来源，脚本不推测 ID、版本或链接。')
    header = source_trace.sheet_header()
    rows = source_trace.sheet_rows(cases_doc)
    if not rows:
        doc.add_paragraph('未记录需求或参考来源。')
    if not source_trace.navigation(cases_doc)['draft_folder_url']:
        doc.add_paragraph('草稿目录：未提供草稿目录链接（不推测 URL）')
    for row in rows:
        paragraph = doc.add_paragraph()
        paragraph.add_run(f'{header[0]}：{row[0]} · {header[1]}：{row[1]}')
        for key, label, value in zip(source_trace.SHEET_COLUMNS[2:], header[2:], row[2:]):
            if value in ('', None):
                continue
            paragraph.add_run().add_break()
            paragraph.add_run(label + '：')
            if key in source_trace.URL_KEYS and source_trace._is_url(value):
                add_docx_hyperlink(paragraph, source_trace.link_text(key, row), value)
            else:
                paragraph.add_run(str(value))


def add_docx_hyperlink(paragraph, text, url):
    """Clickable external link: a w:hyperlink bound to a document relationship."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    props.extend([color, underline])
    run.append(props)
    node = OxmlElement('w:t')
    node.text = text
    node.set(qn('xml:space'), 'preserve')
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def markdown_content(cases_doc, mapping):
    fields = mapping.get('columns', ['用例编号', '用例标题', '前置条件', '测试步骤', '预期结果'])
    if isinstance(fields, dict):
        fields = list(fields.values())
    check_fields(cases_doc['cases'], {str(i): f for i,f in enumerate(fields)})
    lines = ['# ' + cases_doc.get('document', {}).get('title', '测试用例草稿'), '']
    if mapping.get('mode') == 'sections':
        for case in cases_doc['cases']:
            lines += ['## ' + str(case['用例编号']), '']
            for field in fields:
                lines += ['### ' + field, field_text(case, field, mapping), '']
    else:
        def esc(value):
            return text_value(value, mapping).replace('\\', '\\\\').replace('|', '\\|').replace('\n', '<br>')
        header = ['| ' + ' | '.join(fields) + ' |', '| ' + ' | '.join('---' for _ in fields) + ' |']
        block = list(header)
        tables = []
        for case in cases_doc['cases']:
            row = '| ' + ' | '.join(esc(field_text(case, f, mapping)) for f in fields) + ' |'
            try:
                lark_io.split_markdown('\n'.join(block + [row]) + '\n')
            except ValueError:
                if len(block) == len(header):
                    raise ValueError('One case table row exceeds draft block limit; use sections or split fields explicitly')
                tables.append('\n'.join(block))
                block = list(header)
                # A row is never truncated or separated from its header.
                lark_io.split_markdown('\n'.join(block + [row]) + '\n')
            block.append(row)
        tables.append('\n'.join(block))
        lines += ['\n\n'.join(tables)]
    if mapping.get('sources_section', True) is not False and source_trace.has_content(cases_doc):
        lines += ['', source_trace.markdown_section(cases_doc).rstrip('\n')]
    return '\n'.join(lines) + '\n'


def render_markdown(cases_doc, mapping, out_path, overwrite=False):
    out = protect_output(out_path, overwrite=overwrite)
    content = markdown_content(cases_doc, mapping)
    out.write_text(content, encoding='utf-8')
    located = '## 来源与参考' in content and source_trace.has_content(cases_doc)
    return {'out': str(out), 'cases': len(cases_doc['cases']),
            'sources': dict(source_trace.summary(cases_doc), location='appended_section' if located
                            else source_location(mapping, cases_doc, 'sources_section'))}


def default_home():
    return Path(os.environ.get('LARK_TESTCASE_WRITE_HOME', '~/.config/lark-testcase-write')).expanduser() / 'default-template'


def resolve_template(explicit=None):
    if explicit:
        path = Path(explicit)
        return {'template': str(path), 'mapping': str(path.with_suffix('.mapping.json')), 'source': 'explicit'}
    home = default_home()
    manifest = home / 'manifest.json'
    if manifest.exists():
        result = json.loads(manifest.read_text())
        return {'template': str(home/result['template']), 'mapping': str(home/'mapping.json'), 'source': 'user_default'}
    path = Path(__file__).resolve().parents[1]/'assets/templates/标准用例表.xlsx'
    return {'template': str(path), 'mapping': str(path.with_suffix('.mapping.json')), 'source': 'builtin'}


def set_default_template(path, mapping):
    source = Path(path).resolve()
    home = default_home().resolve()
    if home.is_relative_to(source.parent) or source.is_relative_to(home):
        raise ValueError('Default directory overlaps source')
    home.mkdir(parents=True, exist_ok=True)
    target = home/('template'+source.suffix)
    shutil.copy2(source, target)
    (home/'mapping.json').write_text(json.dumps(mapping, ensure_ascii=False, indent=2)+'\n')
    (home/'manifest.json').write_text(json.dumps({'template': target.name})+'\n')
    return resolve_template()
