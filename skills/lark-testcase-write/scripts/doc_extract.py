"""Read original documents mechanically with stable source locations."""
import argparse
import csv
import hashlib
import io
import json
import subprocess
import zipfile
from datetime import date, datetime, time, timedelta
from html.parser import HTMLParser
from pathlib import Path
import lark_io


def scalar(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    # ArrayFormula stores the expression in text; str(object) loses it.
    if type(value).__name__ == 'ArrayFormula' and hasattr(value, 'text'):
        return value.text
    return str(value)


def scalar_fields(value, key='value'):
    fields = {key: scalar(value)}
    if value is not None and not isinstance(value, (str, int, float, bool)):
        fields[key + '_type'] = type(value).__name__
    if type(value).__name__ == 'ArrayFormula' and hasattr(value, 'ref'):
        fields[key + '_ref'] = value.ref
    return fields


def html_span(attributes, key):
    raw = attributes.get(key, '1')
    try:
        value = int(raw)
        if 1 <= value <= 10000:
            return value, None
    except (TypeError, ValueError):
        pass
    return 1, dict(attribute=key, original=raw, used=1)


class HTMLText(HTMLParser):
    """HTML grid locations account for row/column spans without interpreting cells."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks, self.parts, self.table, self.frames = [], [], 0, []
        self.hidden = 0

    def flush(self, location):
        if self.parts:
            self.blocks.append({'location': location, 'text': ''.join(self.parts)})
            self.parts = []

    def finish_cell(self):
        if not self.frames or not self.frames[-1].get('cell'):
            return
        frame = self.frames[-1]; cell = frame['cell']
        self.blocks.append(dict(location=f'表{frame["number"]}!R{frame["row"]}C{cell["column"]}', text=''.join(cell['parts']).rstrip('\n'), rowspan=cell['rowspan'], colspan=cell['colspan'], **({'span_adjustments': cell['span_adjustments']} if cell['span_adjustments'] else {})))
        frame['cell'] = None

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'):
            self.hidden += 1; return
        if self.hidden:
            return
        if tag == 'table':
            self.flush('正文'); self.table += 1
            self.frames.append(dict(number=self.table, row=0, column=0, occupied=set(), cell=None))
        elif self.frames and tag == 'tr':
            self.finish_cell(); self.frames[-1]['row'] += 1; self.frames[-1]['column'] = 0
        elif self.frames and tag in ('td', 'th'):
            self.finish_cell(); frame = self.frames[-1]; attributes = dict(attrs)
            rowspan, row_note = html_span(attributes, 'rowspan')
            colspan, col_note = html_span(attributes, 'colspan')
            column = frame['column'] + 1
            while any((frame['row'], c) in frame['occupied'] for c in range(column, column + colspan)):
                column += 1
            for row in range(frame['row'], frame['row'] + rowspan):
                for col in range(column, column + colspan):
                    frame['occupied'].add((row, col))
            frame['column'] = column + colspan - 1
            frame['cell'] = dict(column=column, rowspan=rowspan, colspan=colspan, parts=[], span_adjustments=[n for n in (row_note, col_note) if n])
        elif tag == 'br':
            self.append_text('\n')

    def append_text(self, value):
        if self.frames and self.frames[-1].get('cell'):
            self.frames[-1]['cell']['parts'].append(value)
        else:
            self.parts.append(value)

    def handle_endtag(self, tag):
        if tag in ('script', 'style'):
            self.hidden = max(0, self.hidden - 1); return
        if self.hidden:
            return
        if tag in ('td', 'th'):
            self.finish_cell()
        elif tag == 'table' and self.frames:
            self.finish_cell(); self.frames.pop(); self.flush('正文')
        elif tag in ('p', 'div', 'h1', 'h2', 'h3'):
            if self.frames and self.frames[-1].get('cell'):
                self.append_text('\n')
            else:
                self.flush('正文')

    def handle_data(self, data):
        if not self.hidden:
            self.append_text(data)


def extract_document(path):
    source = Path(path)
    raw = source.read_bytes()
    result = dict(status='ok', name=source.name, sha256=hashlib.sha256(raw).hexdigest(), size=len(raw), format=None, blocks=[], ranges=[])
    head = raw[:8192].lstrip(b'\xef\xbb\xbf \t\r\n').lower()
    kind = 'unknown'
    if raw.startswith(b'PK'):
        try:
            with zipfile.ZipFile(source) as z:
                names = z.namelist()
            kind = 'xlsx' if 'xl/workbook.xml' in names else 'docx' if 'word/document.xml' in names else 'zip'
        except zipfile.BadZipFile:
            kind = 'invalid_zip'
    elif head.startswith(b'{\\rtf'):
        kind = 'rtf'
    elif raw.startswith(b'\xd0\xcf\x11\xe0'):
        kind = 'doc' if source.suffix.lower() == '.doc' else 'xls'
    elif b'<html' in head or b'<table' in head or head.startswith(b'<!doctype html'):
        kind = 'html'
    elif raw.startswith(b'%PDF'):
        kind = 'pdf'
    elif raw.startswith((b'\x89PNG', b'\xff\xd8\xff', b'GIF8')):
        kind = 'image'
    else:
        try:
            raw.decode('utf-8-sig')
            if b'\x00' not in raw:
                kind = source.suffix.lower().lstrip('.') if source.suffix.lower() in {'.csv', '.tsv', '.md', '.txt', '.json'} else 'txt'
        except UnicodeError:
            pass
    result['format'] = kind
    try:
        if kind == 'xlsx':
            from openpyxl import load_workbook
            with source.open('rb') as f, source.open('rb') as cached:
                book = load_workbook(f, data_only=False)
                values = load_workbook(cached, data_only=True)
                for sheet in book:
                    merges = [str(r) for r in sheet.merged_cells.ranges]
                    result['ranges'].append(dict(sheet=sheet.title, rows=sheet.max_row, columns=sheet.max_column, merged=merges))
                    for row in sheet:
                        cells = []
                        for cell in row:
                            if cell.value is not None:
                                entry = dict(coordinate=cell.coordinate, **scalar_fields(cell.value))
                                if cell.data_type == 'f':
                                    entry.update(scalar_fields(cell.value, 'formula'))
                                    entry.update(scalar_fields(values[sheet.title][cell.coordinate].value, 'cached_value'))
                                cells.append(entry)
                        anchors = []
                        for merged in sheet.merged_cells.ranges:
                            if merged.min_row <= row[0].row <= merged.max_row:
                                anchors.append(dict(range=str(merged), anchor=sheet.cell(merged.min_row, merged.min_col).coordinate, **scalar_fields(sheet.cell(merged.min_row, merged.min_col).value)))
                        if cells or anchors:
                            result['blocks'].append(dict(location=f'{sheet.title}!{row[0].row}', cells=cells, merged_anchors=anchors))
                book.close(); values.close()
        elif kind == 'docx':
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            doc = Document(source)
            paragraph, table = 0, 0
            for element in doc.element.body.iterchildren():
                if element.tag.endswith('}p'):
                    paragraph += 1
                    result['blocks'].append(dict(location=f'段落 {paragraph}', text=Paragraph(element, doc).text))
                elif element.tag.endswith('}tbl'):
                    table += 1
                    for ri, row in enumerate(Table(element, doc).rows, 1):
                        for ci, cell in enumerate(row.cells, 1):
                            result['blocks'].append(dict(location=f'表{table}!R{ri}C{ci}', text=cell.text))
            result['ranges'] = [dict(paragraphs=paragraph, tables=table)]
        elif kind == 'html':
            parser = HTMLText(); parser.feed(raw.decode('utf-8-sig')); parser.flush('正文')
            result['blocks'] = parser.blocks
            result['ranges'] = [dict(blocks=len(parser.blocks), tables=parser.table)]
        elif kind in ('csv', 'tsv'):
            reader = csv.reader(io.StringIO(raw.decode('utf-8-sig')), delimiter='\t' if kind == 'tsv' else ',')
            for ri, row in enumerate(reader, 1):
                result['blocks'].append(dict(location=f'行 {ri}', cells=[dict(coordinate=f'R{ri}C{ci}', value=value) for ci, value in enumerate(row, 1)]))
            result['ranges'] = [dict(rows=len(result['blocks']))]
        elif kind in ('txt', 'md', 'json'):
            result['blocks'] = [dict(location=f'行 {i}', text=line) for i, line in enumerate(raw.decode('utf-8-sig').splitlines(), 1)]
            result['ranges'] = [dict(lines=len(result['blocks']))]
        elif kind in ('doc', 'rtf'):
            r = subprocess.run(['textutil', '-convert', 'txt', '-stdout', str(source)], capture_output=True, timeout=120)
            if r.returncode:
                raise ValueError('textutil could not read this file')
            result['blocks'] = [dict(location=f'行 {i}', text=line) for i, line in enumerate(r.stdout.decode('utf-8').splitlines(), 1)]
            result['ranges'] = [dict(lines=len(result['blocks']))]
        else:
            raise ValueError('unsupported format')
    except ImportError as exc:
        raise ValueError(f'missing dependency: {exc.name}') from exc
    except (OSError, ValueError, UnicodeError, zipfile.BadZipFile, subprocess.SubprocessError) as exc:
        result.update(status='extraction_required', blocks=[], ranges=[], reason=str(exc), suggestion='Convert using drive +import then read the cloud document, or provide a supported readable format.')
    return result


def render_extract(result):
    return '\n\n'.join([f'# {result["name"]}', f'格式: {result["format"]}; SHA-256: {result["sha256"]}; 状态: {result["status"]}', json.dumps(result['ranges'], ensure_ascii=False), *[f'## {b["location"]}\n\n' + (b['text'] if 'text' in b else json.dumps(b, ensure_ascii=False)) for b in result['blocks']], result.get('suggestion', '')]) + '\n'


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('file'); parser.add_argument('--out', required=True); parser.add_argument('--overwrite', action='store_true')
    args = parser.parse_args()
    try:
        directory = lark_io.safe_output(args.out, [args.file], overwrite=True, directory=True)
        outputs = [directory / (Path(args.file).name + suffix) for suffix in ('.extract.json', '.extract.md')]
        for out in outputs:
            lark_io.safe_output(out, [args.file], overwrite=args.overwrite)
        result = extract_document(args.file)
        json_content = json.dumps(result, ensure_ascii=False, indent=2) + '\n'
        markdown_content = render_extract(result)
        directory.mkdir(parents=True, exist_ok=True)
        outputs[0].write_text(json_content)
        outputs[1].write_text(markdown_content)
        print(json.dumps(dict(status=result['status'], out=str(directory), blocks=len(result['blocks'])), ensure_ascii=False))
        return 0 if result['status'] == 'ok' else 2
    except Exception as exc:
        # CLI boundary: dependency/parser/serialization failures still emit JSON.
        print(json.dumps(dict(status='failed', error=str(exc)), ensure_ascii=False)); return 2


if __name__ == '__main__':
    raise SystemExit(main())
