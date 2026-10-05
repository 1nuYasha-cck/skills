"""Offline regressions; fixtures never touch user source files or Lark."""
import copy
import contextlib
import io
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
import lark_io
import templates
import write
from coverage import compute_coverage, compute_mcdc
from doc_extract import extract_document
from openpyxl import Workbook, load_workbook
from docx import Document


def cases():
    return {'document': {'title': '测试草稿'}, 'requirements': [{'id': 'R'}],
        'coverage': {'items': [{'id':'I','requirement':'R','method':'边界值'}]},
        'design_requirements': {'targets': {'requirement':1, 'item':1, 'mcdc':1}},
        'cases': [{'用例编号':'T1','用例标题':'=not_a_formula','需求编号':['R'],'覆盖项':['I'],'设计方法':['边界值'],
                   '前置条件':'初态','测试步骤':['动作1','动作2'],'预期结果':['预期1','预期2'],'优先级':'高'}]}


class TransportTests(unittest.TestCase):
    def test_identity_and_dry_run(self):
        for args in (['--as','bot'], ['--as=bot']):
            with self.assertRaises(ValueError): lark_io.run_lark(args)
        with patch('lark_io.subprocess.run') as run:
            result=lark_io.run_lark(['drive','+upload'],dry_run=True)
            run.assert_not_called()
            self.assertIn('user',result['command'])

    def test_error_and_redaction(self):
        def runner(cmd, **kwargs):
            return subprocess.CompletedProcess(cmd, 1, '', json.dumps({'error':{'code':'permission','message':'denied','access_token':'private'}}))
        with self.assertRaises(lark_io.LarkError) as caught:
            lark_io.run_lark(['base','+record-list'],runner=runner)
        self.assertNotIn('private',json.dumps(caught.exception.as_dict()))

    def test_retry_read_not_write(self):
        with patch('lark_io.time.sleep'):
            for command, attempts in [('+record-list',3),('+record-batch-create',1)]:
                calls=[]
                def runner(cmd, **kwargs):
                    calls.append(cmd)
                    return subprocess.CompletedProcess(cmd,1,'',json.dumps({'error':{'code':'503','message':'temporary'}}))
                with self.assertRaises(lark_io.LarkError): lark_io.run_lark(['base',command],runner=runner)
                self.assertEqual(len(calls),attempts)

    def test_matrix_pagination(self):
        pages=[{'data':{'record_id_list':['r1'],'fields':['Name'],'data':[['one']],'has_more':True}},
               {'data':{'record_id_list':['r2'],'fields':['Name'],'data':[['two']],'has_more':False}}]
        with patch('lark_io.run_lark',side_effect=pages) as run:
            result=lark_io.list_records('base','table',page_size=1)
            self.assertTrue(result['complete'])
            self.assertEqual(result['records'][1]['fields'],{'Name':'two'})
            self.assertIn('1',run.call_args.args[0])

    def test_partial_retains_records(self):
        with patch('lark_io.run_lark',side_effect=[{'records':[{'record_id':'r1','fields':{}}],'has_more':True},lark_io.LarkError('permission','denied')]):
            result=lark_io.list_records('base','table')
            self.assertFalse(result['complete'])
            self.assertEqual(len(result['records']),1)

    def test_bad_matrix(self):
        with patch('lark_io.run_lark',return_value={'data':{'record_id_list':['r'],'fields':['a'],'data':[[]],'has_more':False}}):
            self.assertFalse(lark_io.list_records('b','t')['complete'])

    def test_search_fields(self):
        with patch('lark_io.run_lark',return_value={'records':[],'has_more':False}) as run:
            self.assertTrue(lark_io.search_records('b','t','word',['Title'])['complete'])
            self.assertIn('--search-field',run.call_args.args[0])

    def test_batches(self):
        with patch('lark_io.run_lark',side_effect=[{'record_ids':['r']},lark_io.LarkError('permission','denied')]) as run:
            result=lark_io.batch_create('b','t',[{'Name':str(i)} for i in range(401)])
            self.assertEqual(result['status'],'partial')
            self.assertEqual(len(result['completed']),200)
            self.assertEqual(len(result['failed_slice']),200)
            self.assertEqual(run.call_count,2)
        with patch('lark_io.run_lark',return_value={}) as run:
            lark_io.batch_update('b','t',[{'record_id':str(i),'fields':{}} for i in range(201)])
            self.assertEqual(run.call_count,2)
            body=json.loads(run.call_args.args[0][-1])
            self.assertIn('200',body['update_records'])


class LocalTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.root=Path(self.tmp.name)
        self.src=self.root/'source';self.src.mkdir()
        self.out=self.root/'output';self.out.mkdir()
    def tearDown(self):
        self.tmp.cleanup()
    def test_extract_formats_and_readonly(self):
        w=Workbook();s=w.active;s['A1']='original';s['B2']='=1+2';s.merge_cells('A1:A3')
        w.save(self.src/'book.xlsx')
        d=Document();d.add_paragraph('before');d.add_table(rows=1,cols=1).cell(0,0).text='table';d.add_paragraph('after');d.save(self.src/'doc.docx')
        (self.src/'fake.xlsx').write_text('<html><table><tr><td>cell</td></tr></table></html>')
        (self.src/'text.md').write_text('# 原文\n内容')
        (self.src/'data.csv').write_text('a,b\n1,2')
        for name,kind in [('book.xlsx','xlsx'),('doc.docx','docx'),('fake.xlsx','html'),('text.md','md'),('data.csv','csv')]:
            path=self.src/name;before=path.read_bytes();result=extract_document(path)
            self.assertEqual(result['format'],kind);self.assertEqual(result['status'],'ok')
            self.assertTrue(result['blocks']);self.assertEqual(before,path.read_bytes())
        blocks=extract_document(self.src/'book.xlsx')['blocks']
        self.assertTrue(any(b.get('merged_anchors') for b in blocks))
        self.assertTrue(any(c.get('formula') for b in blocks for c in b.get('cells',[])))
        blocks=extract_document(self.src/'doc.docx')['blocks']
        self.assertEqual([b['text'] for b in blocks],['before','table','after'])
    def test_unsupported(self):
        p=self.src/'file.pdf';p.write_bytes(b'%PDF unparsed')
        self.assertEqual(extract_document(p)['status'],'extraction_required')
    def template(self):
        w=Workbook();s=w.active;s.title='测试用例';s.append(['编号','标题','步骤','预期'])
        s['A2']='example';s['A2'].font=templates.copy.copy(s['A1'].font);s.row_dimensions[2].height=42
        other=w.create_sheet('Other');other['A1']='=1+1';other['B1']='keep'
        p=self.src/'template.xlsx';w.save(p);return p
    def mapping(self):
        return {'sheet':'测试用例','start_row':2,'style_row':2,'columns':{'A':'用例编号','B':'用例标题','C':'测试步骤','D':'预期结果'},'clear_example_rows':[2,2],'row_mode':'per_step','merge_case_columns':['A','B']}
    def test_fill_steps_preserves(self):
        p=self.template();before=p.read_bytes();out=self.out/'cases.xlsx'
        result=templates.fill_xlsx(p,cases(),self.mapping(),out)
        self.assertEqual(result['rows_written'],2)
        w=load_workbook(out);s=w['测试用例']
        self.assertEqual(s['C3'].value,'动作2');self.assertIn('A2:A3',[str(r) for r in s.merged_cells.ranges])
        self.assertEqual(s['B2'].data_type,'s');self.assertEqual(w['Other']['A1'].value,'=1+1')
        self.assertEqual(s.row_dimensions[3].height,42);w.close();self.assertEqual(before,p.read_bytes())
    def test_fill_unknown_long_and_mismatched_steps(self):
        p=self.template()
        for field,value in [('missing',None),('long','x'*32768),('steps',['onlyone'])]:
            c=cases();m=self.mapping()
            if field=='missing':m['columns']['A']='absent'
            elif field=='long':c['cases'][0]['用例标题']=value
            else:c['cases'][0]['预期结果']=value
            with self.assertRaises(ValueError): templates.fill_xlsx(p,c,m,self.out/(field+'.xlsx'))
    def test_protect_source_exists_and_symlink(self):
        p=self.template()
        with self.assertRaises(ValueError):templates.fill_xlsx(p,cases(),self.mapping(),self.src/'new.xlsx')
        self.assertFalse((self.src/'new.xlsx').exists())
        target=self.out/'link';target.symlink_to(self.src,target_is_directory=True)
        with self.assertRaises(ValueError):templates.protect_output(target/'new.xlsx',[p],True)
        existing=self.out/'exists';existing.write_text('keep')
        with self.assertRaises(ValueError):templates.protect_output(existing)
    def test_docx_and_markdown(self):
        d=Document();d.add_paragraph('keep');t=d.add_table(rows=2,cols=2);t.cell(0,0).text='ID';t.cell(1,0).text='style';p=self.src/'template.docx';d.save(p)
        templates.fill_docx(p,cases(),{'table':0,'columns':{'0':'用例编号','1':'用例标题'}},self.out/'cases.docx')
        doc=Document(self.out/'cases.docx');self.assertEqual(doc.paragraphs[0].text,'keep');self.assertEqual(doc.tables[0].rows[-1].cells[0].text,'T1')
        c=cases();c['cases'][0]['用例标题']='a|b\nnext'
        templates.render_markdown(c,{'columns':['用例标题']},self.out/'cases.md')
        self.assertIn('a\\|b<br>next',(self.out/'cases.md').read_text())
    def test_html_template_blank_workbook_cli(self):
        template=self.src/'fake.xlsx';template.write_text('<table><tr><td>编号</td></tr></table>')
        c=self.src/'cases.json';c.write_text(json.dumps(cases(),ensure_ascii=False))
        mapping=self.src/'mapping.json';mapping.write_text(json.dumps({'sheet':'测试用例','start_row':2,'columns':{'A':'用例编号'},'cells':{'A1':'编号'}}))
        self.assertEqual(write.main(['fill','--template',str(template),'--mapping',str(mapping),'--cases',str(c),'--format','xlsx','--out',str(self.out/'html.xlsx')]),0)
        wb=load_workbook(self.out/'html.xlsx');self.assertEqual(wb.active['A2'].value,'T1');wb.close()
    def test_docx_run_style_preserved(self):
        d=Document();t=d.add_table(rows=1,cols=1);t.cell(0,0).paragraphs[0].add_run('sample').bold=True
        p=self.src/'style.docx';d.save(p)
        templates.fill_docx(p,cases(),{'table':0,'columns':{'0':'用例编号'}},self.out/'style.docx')
        new=Document(self.out/'style.docx');self.assertTrue(new.tables[0].rows[-1].cells[0].paragraphs[0].runs[0].bold)
    def test_mcdc_full_model_and_unknown_case(self):
        c=cases();c['cases']=[dict(c['cases'][0],用例编号=f'T{i}') for i in range(1,4)]
        c['coverage']['mcdc']=[self.decision()]
        self.assertTrue(compute_coverage(c)['achieved'])
        c['coverage']['mcdc'][0]['observations'][0]['case']='missing'
        with self.assertRaises(ValueError):compute_coverage(c)

    def test_default_precedence(self):
        p=self.template()
        with patch.dict('os.environ',{'LARK_TESTCASE_WRITE_HOME':str(self.out/'home')}):
            self.assertEqual(templates.resolve_template()['source'],'builtin')
            templates.set_default_template(p,self.mapping())
            self.assertEqual(templates.resolve_template()['source'],'user_default')
            self.assertEqual(templates.resolve_template(p)['source'],'explicit')
    def test_coverage_unknown_and_empty(self):
        report=compute_coverage(cases());self.assertFalse(report['achieved']);self.assertEqual(report['rates']['requirement'],1)
        c=cases();c['cases'][0]['需求编号']=['unknown']
        with self.assertRaises(ValueError):compute_coverage(c)
        self.assertIsNone(compute_coverage({'cases':[]})['rates']['mcdc'])
    def decision(self):
        return {'decision_id':'D','requirement':'R','conditions':['A','B'],'expression':{'and':['A','B']},'observations':[
            {'case':'T1','values':{'A':True,'B':True}}, {'case':'T2','values':{'A':False,'B':True}}, {'case':'T3','values':{'A':True,'B':False}}]}
    def test_mcdc_unique_cause(self):
        d=self.decision();self.assertEqual(compute_mcdc(d)['rate'],1)
        d['observations'][1]['values']={'A':False,'B':False}
        self.assertIn('A',compute_mcdc(d)['missing'])
        d=self.decision();d['observations'][1]['context']='different'
        self.assertIn('A',compute_mcdc(d)['missing'])
    def test_mcdc_invalid(self):
        for changed in ('conditions','values','expression'):
            d=self.decision()
            if changed=='conditions':d['conditions']=['A']
            elif changed=='values':d['observations'][0]['values']['A']=1
            else:d['expression']={'xor':['A','B']}
            with self.assertRaises(ValueError):compute_mcdc(d)
    def test_upload_dryrun_and_partial(self):
        p=self.src/'draft.md';p.write_text('draft')
        with patch('lark_io.subprocess.run') as runner:
            result=write.upload_draft('folder',[p],'# draft',dry_run=True)
            runner.assert_not_called();self.assertTrue(result['draft_only'])
        with patch('lark_io.upload_file',side_effect=[{'file_token':'token'},lark_io.LarkError('permission','denied')]):
            result=write.upload_draft('folder',[p,p])
            self.assertEqual(result['status'],'partial');self.assertEqual(len(result['files']),1)
    def test_upload_readback(self):
        p=self.src/'draft.md';p.write_text('draft')
        with patch('lark_io.upload_file',return_value={'file_token':'token'}),patch('lark_io.run_lark',return_value={'data':{'files':[{'token':'token'}],'has_more':False}}):
            self.assertTrue(write.upload_draft('folder',[p])['files'][0]['readback'])
    def test_cli_output_protection_and_coverage_target(self):
        p=self.src/'cases.json';p.write_text(json.dumps(cases(),ensure_ascii=False))
        self.assertEqual(write.main(['coverage','--cases',str(p),'--out',str(self.src/'coverage.json')]),2)
        self.assertEqual(write.main(['coverage','--cases',str(p),'--out',str(self.out/'coverage.json')]),0)
        self.assertFalse(json.loads((self.out/'coverage.json').read_text())['achieved'])

    def test_config_explicit_library_and_destination(self):
        config={'config_version':'2.0','libraries':[
            {'name':'one','folders':{'root':'one-root','drafts':'one-draft','bodies':'one-body','originals':'one-original'}},
            {'name':'two','folders':{'root':'two-root','drafts':'two-draft','bodies':'two-body'}}]}
        p=self.src/'draft.md';p.write_text('draft')
        with patch('lark_io.subprocess.run') as runner:
            result=write.upload_draft(config,[p],library='two',folder_key='drafts',dry_run=True)
            self.assertIn('two-draft',result['files'][0]['response']['command'])
            runner.assert_not_called()
            for library,key in [(None,'drafts'),('one',None),('missing','root'),('one','bodies'),('one','originals'),('one','missing')]:
                with self.assertRaises(ValueError):write.upload_draft(config,[p],library=library,folder_key=key,dry_run=True)
            config['libraries'][0]['folders']['alias']='one-body'
            with self.assertRaises(ValueError):write.upload_draft(config,[p],library='one',folder_key='alias',dry_run=True)
            config['libraries'].append(config['libraries'][0])
            with self.assertRaises(ValueError):write.upload_draft(config,[p],library='one',folder_key='drafts',dry_run=True)
    def test_config_cli_and_no_default(self):
        config=self.src/'config.json';config.write_text(json.dumps({'libraries':[{'name':'one','folders':{'root':'root','drafts':'drafts'}}]}))
        draft=self.src/'draft.md';draft.write_text('draft')
        base=['upload','--config',str(config),'--file',str(draft),'--dry-run']
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(write.main(base+['--out',str(self.out/'bad.json')]),2)
            self.assertEqual(write.main(base+['--library','one','--folder-key','drafts','--out',str(self.out/'good.json')]),0)
        report=json.loads((self.out/'good.json').read_text());self.assertIn('drafts',report['files'][0]['response']['command'])
    def test_fill_rejects_merged_and_nonempty_without_output(self):
        for merged in (False,True):
            p=self.template();w=load_workbook(p);ws=w['测试用例']
            ws['A2']=None;ws['C3']='preserve'
            if merged:ws.merge_cells('C3:D3')
            w.save(p);w.close();before=p.read_bytes();m=self.mapping()
            with self.assertRaisesRegex(ValueError,'C3'):
                templates.fill_xlsx(p,cases(),m,self.out/'conflict.xlsx')
            self.assertFalse((self.out/'conflict.xlsx').exists());self.assertEqual(p.read_bytes(),before)
    def test_cli_merged_conflict_json_exit2(self):
        p=self.template();w=load_workbook(p);w['测试用例'].merge_cells('A2:B2');w.save(p);w.close()
        c=self.src/'cases.json';c.write_text(json.dumps(cases(),ensure_ascii=False))
        m=self.mapping();m.pop('clear_example_rows');mapping=self.src/'mapping.json';mapping.write_text(json.dumps(m))
        stdout=io.StringIO()
        with contextlib.redirect_stdout(stdout):
            rc=write.main(['fill','--template',str(p),'--mapping',str(mapping),'--cases',str(c),'--out',str(self.out/'conflict.xlsx')])
        self.assertEqual(rc,2);result=json.loads(stdout.getvalue());self.assertIn('A2:B2',result['error'])
        self.assertNotIn('Traceback',stdout.getvalue());self.assertFalse((self.out/'conflict.xlsx').exists())
    def test_blank_workbook_explicit_formatting(self):
        mapping={'sheet':'测试用例','start_row':2,'columns':{'A':'用例编号','B':'测试步骤'},
                 'column_widths':{'A':20,'B':65},'alignment':{'wrap_text':True,'vertical':'top','horizontal':'left'},'row_height':90,'cells':{'A1':'标题'}}
        out=self.out/'formatted.xlsx';templates.fill_xlsx(None,cases(),mapping,out)
        wb=load_workbook(out);ws=wb.active
        self.assertEqual(ws.column_dimensions['B'].width,65);self.assertEqual(ws.row_dimensions[2].height,90)
        for key in ('A1','A2','B2'):
            self.assertTrue(ws[key].alignment.wrap_text);self.assertEqual(ws[key].alignment.vertical,'top')
        wb.close()
        for key,value in [('row_height',-1),('column_widths',{'A':0}),('alignment',{'vertical':'invalid'})]:
            with self.assertRaises(ValueError):templates.fill_xlsx(None,cases(),dict(mapping,**{key:value}),self.out/'invalid.xlsx')
    def test_blank_default_wrap_and_existing_style_override(self):
        mapping={'sheet':'测试用例','start_row':2,'columns':{'A':'测试步骤'}}
        out=self.out/'default.xlsx';templates.fill_xlsx(None,cases(),mapping,out)
        wb=load_workbook(out);self.assertTrue(wb.active['A2'].alignment.wrap_text);self.assertEqual(wb.active['A2'].alignment.vertical,'top');wb.close()
        p=self.template();m=self.mapping();m.update(alignment={'wrap_text':True,'vertical':'top'},row_height=75)
        out=self.out/'override.xlsx';templates.fill_xlsx(p,cases(),m,out);wb=load_workbook(out)
        self.assertEqual(wb['测试用例']['C3'].alignment.vertical,'top');self.assertEqual(wb['测试用例'].row_dimensions[3].height,75);wb.close()
    def test_docx_paragraph_in_place_and_missing_field(self):
        d=Document();d.add_paragraph('标题行');d.add_paragraph().add_run('用例 {{用例编号}}：{{用例标题}}').bold=True;d.add_paragraph('结尾说明')
        p=self.src/'para.docx';d.save(p);before=p.read_bytes()
        c=cases();c['cases']=[dict(c['cases'][0],用例编号='TC-1',用例标题='甲'),dict(c['cases'][0],用例编号='TC-2',用例标题='乙')]
        out=self.out/'para.docx';templates.fill_docx(p,c,{'paragraphs':[1]},out)
        result=Document(out);self.assertEqual([x.text for x in result.paragraphs],['标题行','用例 TC-1：甲','用例 TC-2：乙','结尾说明'])
        self.assertTrue(result.paragraphs[1].runs[0].bold);self.assertEqual(p.read_bytes(),before)
        del c['cases'][1]['用例标题']
        with self.assertRaisesRegex(ValueError,'用例标题'):templates.fill_docx(p,c,{'paragraphs':[1]},self.out/'missing.docx')
        self.assertFalse((self.out/'missing.docx').exists())
    def test_docx_block_and_split_placeholder(self):
        d=Document();d.add_paragraph('head');d.add_paragraph('ID {{用例编号}}');d.add_paragraph('Title {{用例标题}}');d.add_paragraph('tail')
        p=self.src/'block.docx';d.save(p)
        c=cases();c['cases']=[dict(c['cases'][0],用例编号='1'),dict(c['cases'][0],用例编号='2')]
        templates.fill_docx(p,c,{'paragraphs':[1,2]},self.out/'block.docx')
        texts=[x.text for x in Document(self.out/'block.docx').paragraphs]
        self.assertEqual(texts,['head','ID 1','Title =not_a_formula','ID 2','Title =not_a_formula','tail'])
        with self.assertRaises(ValueError):templates.fill_docx(p,c,{'paragraphs':[1,3]},self.out/'noncontiguous.docx')
        d=Document();para=d.add_paragraph();para.add_run('{{用例');para.add_run('编号}}');d.save(p)
        with self.assertRaisesRegex(ValueError,'spans'):templates.fill_docx(p,c,{'paragraphs':[0]},self.out/'split.docx')
    def test_local_extraction_failure_exit2_both_entries(self):
        p=self.src/'unsupported.pdf';p.write_bytes(b'%PDF unsupported')
        stdout=io.StringIO()
        with contextlib.redirect_stdout(stdout):
            self.assertEqual(write.main(['extract','--input',str(p),'--out',str(self.out/'extract')]),2)
        self.assertEqual(json.loads(stdout.getvalue())['status'],'partial')
        proc=subprocess.run([sys.executable,str(Path(write.__file__).parent/'doc_extract.py'),str(p),'--out',str(self.out/'direct')],capture_output=True,text=True)
        self.assertEqual(proc.returncode,2);self.assertEqual(json.loads(proc.stdout)['status'],'extraction_required')
    def test_cli_unexpected_library_exception_contract(self):
        p=self.template();stdout=io.StringIO()
        with patch('write.inspect_template',side_effect=AttributeError('format issue')),contextlib.redirect_stdout(stdout):
            self.assertEqual(write.main(['inspect-template','--template',str(p),'--out',str(self.out/'error.json')]),2)
        self.assertEqual(json.loads(stdout.getvalue())['status'],'failed')

    def test_canonical_html_spans_and_invalid_span(self):
        p=self.src/'span.xlsx'
        p.write_text('<table><tr><td rowspan="2" colspan="2">first</td><td>third</td></tr><tr><td>lower</td></tr></table>')
        result=extract_document(p)
        blocks={b['text']:b for b in result['blocks'] if b.get('text')}
        self.assertEqual(blocks['lower']['location'],'表1!R2C3')
        self.assertEqual(blocks['first']['rowspan'],2);self.assertEqual(blocks['first']['colspan'],2)
        p.write_text('<table><tr><td rowspan="invalid">first</td><td>second</td></tr></table>')
        result=extract_document(p);self.assertEqual(result['status'],'ok')
        self.assertTrue(result['blocks'][0]['span_adjustments'])
    def test_canonical_time_duration_array_cli(self):
        from datetime import time, timedelta
        from openpyxl.worksheet.formula import ArrayFormula
        p=self.src/'time.xlsx';wb=Workbook();ws=wb.active
        ws['A1']=time(12,34,56);ws['B1']=timedelta(hours=27,seconds=5);ws['B1'].number_format='[h]:mm:ss'
        ws['C1']=ArrayFormula(ref='C1:C2',text='=A1:A2*2');ws['D1']='ordinary'
        wb.save(p);before=p.read_bytes()
        for entry in ('write', 'direct'):
            out=self.out/entry
            command=([sys.executable,str(Path(write.__file__)),'extract','--input',str(p),'--out',str(out)] if entry=='write' else
                     [sys.executable,str(Path(write.__file__).parent/'doc_extract.py'),str(p),'--out',str(out)])
            proc=subprocess.run(command,capture_output=True,text=True)
            self.assertEqual(proc.returncode,0,proc.stdout+proc.stderr)
            result=json.loads((out/'time.xlsx.extract.json').read_text())
            values={c['coordinate']:c for b in result['blocks'] for c in b.get('cells',[])}
            self.assertEqual(values['A1']['value'],'12:34:56');self.assertEqual(values['A1']['value_type'],'time')
            self.assertEqual(values['B1']['value_type'],'timedelta');self.assertEqual(values['C1']['value_type'],'ArrayFormula')
            self.assertEqual(values['C1']['value_ref'],'C1:C2');self.assertEqual(values['C1']['formula'],'=A1:A2*2')
            self.assertIsNone(values['C1']['cached_value']);self.assertEqual(values['D1']['value'],'ordinary')
        self.assertEqual(p.read_bytes(),before)
    def test_canonical_duplicate_page_rejected_as_whole(self):
        first={'records':[{'record_id':'r1','fields':{}}],'has_more':True}
        for second in ({'records':[{'record_id':'r2','fields':{}},{'record_id':'r2','fields':{}}],'has_more':False},
                       {'records':[{'record_id':'r2','fields':{}},{'record_id':'r1','fields':{}}],'has_more':False}):
            with patch('lark_io.run_lark',side_effect=[first,second]):
                result=lark_io.list_records('base','table')
            self.assertFalse(result['complete']);self.assertEqual(result['error']['code'],'pagination')
            self.assertEqual([r['record_id'] for r in result['records']],['r1'])

    def test_table_draft_splits_at_byte_and_line_limits(self):
        from templates import markdown_content
        for count,payload in [(61,'长步骤内容'*70),(170,'short')]:
            c=cases();c['cases']=[dict(c['cases'][0],用例编号=f'CASE-{i:03}',测试步骤=[payload]) for i in range(count)]
            markdown=markdown_content(c,{'columns':['用例编号','用例标题','测试步骤']})
            chunks=lark_io.split_markdown(markdown)
            tables=[block for block in markdown.split('\n\n') if block.startswith('|')]
            self.assertGreater(len(tables),1)
            for table in tables:
                self.assertTrue(table.startswith('| 用例编号 |'));self.assertLessEqual(len(table.encode('utf-8')),6000);self.assertLessEqual(len(table.splitlines()),80)
            for i in range(count):self.assertEqual(markdown.count(f'CASE-{i:03}'),1)
            with patch('lark_io.subprocess.run') as runner:
                result=write.upload_draft('folder',[],markdown,title='Large draft',dry_run=True)
                runner.assert_not_called()
                self.assertEqual(result['status'],'ok');self.assertEqual(result['document']['action'],'create_planned')
    def test_table_single_row_too_large_rejected(self):
        c=cases();c['cases'][0]['用例标题']='文'*3000
        with self.assertRaisesRegex(ValueError,'One case'):
            templates.markdown_content(c,{'columns':['用例编号','用例标题']})
    def test_draft_collision_preflight_preserves_existing(self):
        p=self.src/'upload.txt';p.write_text('upload')
        existing={'token':'old','url':'old-url','name':'Title','type':'docx'}
        with patch('lark_io.list_folder',return_value=[existing]),patch('lark_io.create_doc') as create,patch('lark_io.upload_file') as upload:
            with self.assertRaisesRegex(ValueError,'already exists'):
                write.upload_draft('folder',[p],'# Other\n\nbody',title='Title')
            create.assert_not_called();upload.assert_not_called()
    def test_draft_replacement_and_confirmed_title(self):
        entry={'token':'doc','url':'doc-url','name':'Title','type':'docx'}
        for replace in (False,True):
            preflight=[entry] if replace else []
            with patch('lark_io.list_folder',side_effect=[preflight,[entry]]),patch('lark_io.create_doc',return_value={'document':{'token':'doc','url':'doc-url'}}) as create,patch('lark_io.fetch_doc',return_value={'content':'# Title\nbody'}):
                result=write.upload_draft('folder',[],'# Other\n\nbody',title='Title',replace_draft=replace)
                self.assertEqual(result['document']['action'],'replaced' if replace else 'created')
                self.assertEqual(result['document']['actual_title'],'Title')
                self.assertTrue(create.call_args.args[1].startswith('# Title\n'))
    def test_title_injected_and_wrong_remote_title_detected(self):
        with patch('lark_io.subprocess.run') as runner:
            report=write.upload_draft('folder',[],'body',title='Requested',dry_run=True)
            runner.assert_not_called();self.assertTrue(report['document']['response']['markdown'].startswith('# Requested\n\n'))
            self.assertEqual(report['document']['collision_check'],'not_executed_dry_run')
        with patch('lark_io.list_folder',side_effect=[[],[{'token':'doc','name':'Wrong','type':'docx'}]]),patch('lark_io.create_doc',return_value={'document':{'token':'doc'}}),patch('lark_io.fetch_doc',return_value={'content':'body'}):
            result=write.upload_draft('folder',[],'# Other',title='Requested')
            self.assertEqual(result['status'],'partial');self.assertEqual(result['error']['code'],'draft_title_readback')
    def test_replace_cli_explicit_and_duplicate_titles_rejected(self):
        p=self.src/'draft.md';p.write_text('# Other\n\nbody')
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(write.main(['upload','--folder','folder','--markdown',str(p),'--title','Requested','--replace-draft','--dry-run','--out',str(self.out/'replace.json')]),0)
        self.assertEqual(json.loads((self.out/'replace.json').read_text())['document']['action'],'replace_allowed')
        with patch('lark_io.list_folder',return_value=[{'type':'docx','name':'Title'},{'type':'docx','name':'Title'}]),patch('lark_io.create_doc') as create:
            with self.assertRaisesRegex(ValueError,'Multiple'):
                write.upload_draft('folder',[],'body',title='Title',replace_draft=True)
            create.assert_not_called()
    def test_oversize_upload_preflight_prevents_file_write(self):
        p=self.src/'upload.txt';p.write_text('file')
        with patch('lark_io.upload_file') as upload,patch('lark_io.list_folder') as listing:
            with self.assertRaises(ValueError):write.upload_draft('folder',[p],'文'*3000)
            upload.assert_not_called();listing.assert_not_called()

    def test_extract_existing_empty_directory_and_reject_nonempty(self):
        p=self.src/'req.md';p.write_text('requirement')
        directory=self.out/'extract';directory.mkdir()
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(write.main(['extract','--input',str(p),'--out',str(directory)]),0)
        self.assertTrue((directory/'req.md.extract.json').exists())
        before={f.name:f.read_bytes() for f in directory.iterdir()}
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(write.main(['extract','--input',str(p),'--out',str(directory)]),2)
            self.assertEqual(write.main(['extract','--input',str(p),'--out',str(directory),'--overwrite']),2)
        self.assertEqual(before,{f.name:f.read_bytes() for f in directory.iterdir()})
    def test_coverage_markdown_readable_tables_and_gaps(self):
        from coverage import coverage_markdown
        c=cases();c['requirements'].append({'id':'R2'})
        c['coverage']['items'].append({'id':'I2','requirement':'R2','method':'等价类','description':'未覆盖项'})
        d=self.decision();d['observations']=d['observations'][:1];c['coverage']['mcdc']=[d]
        report=compute_coverage(c);text=coverage_markdown(c,report)
        self.assertNotIn('```json',text);self.assertIn('| R2 | 0 | 缺口 |',text)
        self.assertIn('### 等价类',text);self.assertIn('| D | R | A | — | 缺口 |',text)
        c['cases']=[dict(c['cases'][0],用例编号=f'T{i}') for i in range(1,4)];c['coverage']['mcdc']=[self.decision()]
        text=coverage_markdown(c,compute_coverage(c));self.assertIn('T1 ↔ T2',text)
        empty=coverage_markdown({'cases':[]},compute_coverage({'cases':[]}));self.assertIn('无模型/不适用',empty);self.assertIn('未指定目标',empty)
    def test_coverage_markdown_unmodeled_method_and_unknown_targets(self):
        from coverage import coverage_markdown
        c=cases();c['design_requirements']['targets']={'边界值':1,'等价类':1,'requrement':1}
        report=compute_coverage(c);before=copy.deepcopy(report)
        text=coverage_markdown(c,report)
        self.assertIn('| 边界值 | 100.0% | 100.0% | 达到 |',text)
        for key in ('等价类','requrement'):
            self.assertFalse(report['targets_reached'][key])
            self.assertIn(f'| {key} | 无模型/不适用 | 100.0% | 未达到 |',text)
        self.assertIn('- 未建模的目标：等价类, requrement',text)
        self.assertIn('目标结论：存在未达到目标',text)
        self.assertFalse(report['achieved']);self.assertEqual(report,before)
        self.assertEqual(text.count('| 等价类 |'),1)
    def test_coverage_markdown_null_rate_target_and_modeled_gap(self):
        from coverage import coverage_markdown
        c=cases();c['design_requirements']['targets']={'mcdc':0,'边界值':1}
        c['cases'][0]['覆盖项']=[]
        report=compute_coverage(c);text=coverage_markdown(c,report)
        self.assertIn('| mcdc | 无模型/不适用 | 0.0% | 未达到 |',text)
        self.assertIn('| 边界值 | 0.0% | 100.0% | 未达到 |',text)
        self.assertIn('- 未建模的目标：mcdc\n',text)
        self.assertIn('- 覆盖项：I',text)
        c['design_requirements']['targets']={'边界值':0}
        text=coverage_markdown(c,compute_coverage(c))
        self.assertIn('- 未建模的目标：无',text)
        self.assertIn('目标结论：全部达到',text)
    def test_fill_compact_summary_and_full_sidecar(self):
        p=self.template();c=self.src/'cases.json';c.write_text(json.dumps(cases(),ensure_ascii=False))
        m=self.src/'mapping.json';m.write_text(json.dumps(self.mapping(),ensure_ascii=False))
        out=self.out/'filled.xlsx';stdout=io.StringIO()
        with contextlib.redirect_stdout(stdout):
            self.assertEqual(write.main(['fill','--template',str(p),'--mapping',str(m),'--cases',str(c),'--out',str(out)]),0)
        summary=json.loads(stdout.getvalue());self.assertNotIn('merges',summary);self.assertNotIn('sha256',summary)
        self.assertEqual(summary['rows_written'],2);self.assertEqual(summary['merge_count'],2)
        detail=json.loads(Path(summary['result_out']).read_text());self.assertEqual(detail['merges'],['A2:A3','B2:B3'])
    def test_sheet_rename_preserves_other_sheets_and_source(self):
        p=self.template();before=p.read_bytes();m=self.mapping();m['output_sheet']='功能用例'
        out=self.out/'renamed.xlsx';templates.fill_xlsx(p,cases(),m,out)
        wb=load_workbook(out);self.assertEqual(wb.sheetnames,['功能用例','Other']);self.assertEqual(wb['Other']['A1'].value,'=1+1');wb.close();self.assertEqual(p.read_bytes(),before)
        m['output_sheet']='Other'
        with self.assertRaisesRegex(ValueError,'already exists'):templates.fill_xlsx(p,cases(),m,self.out/'collision.xlsx')
    def test_sheet_creation_and_collision(self):
        p=self.template();before=p.read_bytes();m=self.mapping();m.update(sheet='新用例表',create_sheet=True)
        out=self.out/'created.xlsx';templates.fill_xlsx(p,cases(),m,out)
        wb=load_workbook(out);self.assertEqual(wb['测试用例']['A2'].value,'example');self.assertEqual(wb['新用例表']['C3'].value,'动作2');self.assertTrue(wb['新用例表']['C3'].alignment.wrap_text);wb.close();self.assertEqual(p.read_bytes(),before)
        m['sheet']='Other'
        with self.assertRaisesRegex(ValueError,'already exists'):templates.fill_xlsx(p,cases(),m,self.out/'bad.xlsx')
    def test_builtin_mapping_width_and_long_id(self):
        for name in ('标准用例表','步骤展开表'):
            directory=Path(write.__file__).resolve().parents[1]/'assets/templates'
            mapping=json.loads((directory/(name+'.mapping.json')).read_text())
            self.assertGreaterEqual(mapping['column_widths']['A'],28)
            c=cases();c['cases'][0]['用例编号']='TC_THERM_001_01_14V'
            out=self.out/(name+'.xlsx');templates.fill_xlsx(directory/(name+'.xlsx'),c,mapping,out)
            wb=load_workbook(out);self.assertEqual(wb.active['A2'].value,'TC_THERM_001_01_14V');self.assertEqual(wb.active.column_dimensions['A'].width,30);wb.close()


if __name__=='__main__':unittest.main()
