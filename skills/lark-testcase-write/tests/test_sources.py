"""Structured sources, navigation and backlink guards; offline, fake Lark only."""
import contextlib
import copy
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
import lark_io
import source_trace
import templates
import write
from coverage import compute_coverage
from openpyxl import load_workbook
from docx import Document

REQ_URL = 'https://example.invalid/docx/ReqDoc0001#section-3'
INDEX_URL = 'https://example.invalid/base/FormalBase01?table=tblCases0001&record=recCase00001'
BODY_URL = 'https://example.invalid/docx/BodyDoc0001#blockAnchor01'
DRAFT_PEER = 'https://example.invalid/docx/PeerDraft001'
FOLDER_URL = 'https://example.invalid/drive/folder/DraftFolder01'


def plain_cases():
    return {'document': {'title': '测试草稿'}, 'requirements': [{'id': 'R'}],
            'coverage': {'items': [{'id': 'I', 'requirement': 'R', 'method': '边界值'}]},
            'design_requirements': {'targets': {'requirement': 1, 'item': 1}},
            'cases': [{'用例编号': 'T1', '用例标题': '标题', '需求编号': ['R'], '覆盖项': ['I'], '设计方法': ['边界值'],
                       '前置条件': '初态', '测试步骤': ['动作1', '动作2'], '预期结果': ['预期1', '预期2'], '优先级': '高'}]}


def traced_cases():
    doc = plain_cases()
    doc['document']['sources'] = [{'kind': 'requirement', 'title': '需求规格', 'url': REQ_URL, 'section': '3.2 速度控制'}]
    doc['document']['navigation'] = {'editable_drafts': [DRAFT_PEER]}
    doc['cases'][0]['参考来源'] = [{'kind': 'reference_case', 'type': '修改后复用', 'case_no': 'TC_001_002',
                                   'system_id': 'recCase00001', 'index_url': INDEX_URL, 'body_url': BODY_URL,
                                   'reason': '阈值改为 9V，步骤保留', 'unconfirmed': ['回差待需求确认']},
                                  '需求文档!A2']
    return doc


def docx_links(path):
    """(text, target) of every w:hyperlink, resolved through word/_rels; fails on dangling or internal ids."""
    import zipfile
    from xml.etree import ElementTree
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
          'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
          'rel': 'http://schemas.openxmlformats.org/package/2006/relationships'}
    with zipfile.ZipFile(path) as archive:
        body = ElementTree.fromstring(archive.read('word/document.xml'))
        rels = ElementTree.fromstring(archive.read('word/_rels/document.xml.rels'))
    targets = {}
    for rel in rels.findall('rel:Relationship', ns):
        if rel.get('Type').endswith('/hyperlink'):
            assert rel.get('TargetMode') == 'External', rel.attrib
            targets[rel.get('Id')] = rel.get('Target')
    links = []
    for link in body.iter('{%s}hyperlink' % ns['w']):
        r_id = link.get('{%s}id' % ns['r'])
        assert r_id in targets, r_id
        links.append((''.join(t.text or '' for t in link.iter('{%s}t' % ns['w'])), targets[r_id]))
    return links


class SourceRenderTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.src = self.root/'source'; self.src.mkdir()
        self.out = self.root/'output'; self.out.mkdir()
        self.home = patch.dict(os.environ, {'LARK_TESTCASE_WRITE_HOME': str(self.root/'home')})
        self.home.start()

    def tearDown(self):
        self.home.stop()
        self.tmp.cleanup()

    def builtin(self):
        resolved = templates.resolve_template()
        self.assertEqual(resolved['source'], 'builtin')
        return resolved['template'], json.loads(Path(resolved['mapping']).read_text())

    def test_markdown_keeps_links_readable_ids_and_missing_marks(self):
        out = self.out/'cases.md'
        report = templates.render_markdown(traced_cases(), {'mode': 'table'}, out)
        text = out.read_text()
        for expected in (REQ_URL, '3.2 速度控制', 'TC_001_002', INDEX_URL.replace('(', '%28'), BODY_URL,
                         '阈值改为 9V，步骤保留', '回差待需求确认', '系统ID（追溯）：recCase00001',
                         '内容版本：未提供', '未提供草稿目录链接（不推测 URL）', '引用：需求文档!A2'):
            self.assertIn(expected, text)
        self.assertEqual(report['sources']['location'], 'appended_section')
        self.assertEqual(report['sources']['missing'][0]['fields'], ['content_version'])
        for chunk in lark_io.split_markdown(text):
            self.assertLessEqual(len(chunk.encode()), 6000)

    def test_xlsx_builtin_template_keeps_main_sheet_and_adds_source_sheet(self):
        template, mapping = self.builtin()
        before = Path(template).read_bytes()
        templates.fill_xlsx(template, plain_cases(), mapping, self.out/'plain.xlsx')
        report = templates.fill_xlsx(template, traced_cases(), mapping, self.out/'traced.xlsx')
        self.assertEqual(before, Path(template).read_bytes())
        plain, traced = load_workbook(self.out/'plain.xlsx'), load_workbook(self.out/'traced.xlsx')
        self.assertEqual(plain.sheetnames, load_workbook(template).sheetnames)
        self.assertEqual(traced.sheetnames, plain.sheetnames + ['来源与参考'])
        self.assertEqual([[c.value for c in r] for r in plain['测试用例'].iter_rows()],
                         [[c.value for c in r] for r in traced['测试用例'].iter_rows()])
        sheet = traced['来源与参考']
        values = [[c.value for c in r] for r in sheet.iter_rows()]
        self.assertEqual(values[0][:2], ['范围', '序号'])
        flat = [v for row in values for v in row if v]
        for expected in (REQ_URL, 'TC_001_002', INDEX_URL, BODY_URL, '阈值改为 9V，步骤保留', 'recCase00001', '需求文档!A2'):
            self.assertIn(expected, flat)
        links = {c.hyperlink.target for row in sheet.iter_rows() for c in row if c.hyperlink}
        self.assertTrue({REQ_URL, INDEX_URL, BODY_URL} <= links)
        self.assertEqual(report['sources']['location'], '来源与参考')

    def test_docx_and_mapped_column_render_readable_not_json(self):
        template = self.src/'t.docx'
        d = Document(); t = d.add_table(rows=1, cols=2); d.save(template)
        templates.fill_docx(template, traced_cases(), {'table': 0, 'style_row': 0, 'columns': {'0': '用例编号', '1': '参考来源'}},
                            self.out/'cases.docx')
        doc = Document(self.out/'cases.docx')
        cell = doc.tables[0].rows[-1].cells[1].text
        self.assertIn('可读编号：TC_001_002', cell)
        self.assertIn('适配理由：阈值改为 9V，步骤保留', cell)
        self.assertNotIn('{"', cell)
        body = '\n'.join(p.text for p in doc.paragraphs)
        for expected in ('来源与参考', '链接：需求规格', '索引记录链接：索引记录', '正文章节链接：正文章节', '回差待需求确认'):
            self.assertIn(expected, body)
        self.assertEqual({target for _, target in docx_links(self.out/'cases.docx')}, {REQ_URL, INDEX_URL, BODY_URL})

    def test_pure_path_and_structured_sources_keep_path_in_xlsx_and_docx(self):
        doc = traced_cases()
        doc['document']['sources'].insert(0, {'path': '需求/本地需求说明.docx'})
        doc['cases'][0]['参考来源'].append({'kind': 'search_report', 'path': '检索/检索报告.md', 'note': '已复核'})
        doc['document']['navigation']['draft_folder_url'] = FOLDER_URL
        header, rows = source_trace.sheet_header(), source_trace.sheet_rows(doc)
        self.assertEqual(len(header), len(source_trace.SHEET_COLUMNS))
        self.assertTrue(all(len(row) == len(header) for row in rows))
        self.assertEqual(set(source_trace.COLUMN_WIDTHS), set(source_trace.SHEET_COLUMNS))
        path_col = source_trace.SHEET_COLUMNS.index('path')
        self.assertEqual(header[path_col], '本地路径')
        self.assertEqual([row[path_col] for row in rows if row[path_col]], ['需求/本地需求说明.docx', '检索/检索报告.md'])
        self.assertTrue(all(row[path_col] == '' for row in rows if row[0] == '导航'))
        self.assertIn('本地路径：检索/检索报告.md', source_trace.cell_text(doc['cases'][0]['参考来源']))

        template, mapping = self.builtin()
        templates.fill_xlsx(template, doc, mapping, self.out/'mixed.xlsx')
        sheet = load_workbook(self.out/'mixed.xlsx')['来源与参考']
        values = [[c.value for c in r] for r in sheet.iter_rows()]
        self.assertEqual(values[0], header)
        column = {key: i for i, key in enumerate(source_trace.SHEET_COLUMNS)}
        self.assertEqual([row[path_col] for row in values[1:] if row[path_col]], ['需求/本地需求说明.docx', '检索/检索报告.md'])
        self.assertEqual(sheet.column_dimensions['O'].width, source_trace.COLUMN_WIDTHS['path'])
        self.assertEqual(sheet.column_dimensions['R'].width, source_trace.COLUMN_WIDTHS['其他字段'])
        for row in sheet.iter_rows(min_row=2):
            for i, cell in enumerate(row):
                if cell.hyperlink:
                    self.assertIn(source_trace.SHEET_COLUMNS[i], source_trace.URL_KEYS)
                    self.assertEqual(cell.hyperlink.target, cell.value)
        links = {c.hyperlink.target for row in sheet.iter_rows() for c in row if c.hyperlink}
        self.assertEqual(links, {FOLDER_URL, REQ_URL, INDEX_URL, BODY_URL})
        self.assertEqual(values[1][column['url']], FOLDER_URL)

        template_docx = self.src/'t.docx'
        d = Document(); d.add_table(rows=1, cols=1); d.save(template_docx)
        templates.fill_docx(template_docx, doc, {'table': 0, 'style_row': 0, 'columns': {'0': '用例编号'}}, self.out/'mixed.docx')
        body = '\n'.join(p.text for p in Document(self.out/'mixed.docx').paragraphs)
        self.assertIn('本地路径：需求/本地需求说明.docx', body)
        self.assertIn('本地路径：检索/检索报告.md', body)

    def test_docx_source_links_are_relationship_hyperlinks(self):
        doc = traced_cases()
        nav_url = 'https://example.invalid/wiki/CatalogPage01'
        doc['document']['navigation'].update(draft_folder_url=FOLDER_URL, links=[{'title': '用例目录', 'url': nav_url}])
        template = self.src/'t.docx'
        d = Document(); d.add_table(rows=1, cols=1); d.save(template)
        templates.fill_docx(template, doc, {'table': 0, 'style_row': 0, 'columns': {'0': '用例编号'}}, self.out/'linked.docx')
        links = docx_links(self.out/'linked.docx')
        self.assertEqual(sorted(links), sorted([('打开草稿目录', FOLDER_URL), ('用例目录', nav_url), ('需求规格', REQ_URL),
                                                ('索引记录', INDEX_URL), ('正文章节', BODY_URL)]))
        body = '\n'.join(p.text for p in Document(self.out/'linked.docx').paragraphs)
        self.assertNotIn('未提供草稿目录链接', body)
        for url in (FOLDER_URL, nav_url, REQ_URL, INDEX_URL, BODY_URL):
            self.assertNotIn(url, body)

    def test_unknown_ref_and_extra_fields_preserved_invalid_link_rejected(self):
        doc = plain_cases()
        doc['cases'][0]['参考来源'] = {'type': '新写', 'ref': 'Legacy-REF?42', 'note': '依据R', 'custom_tag': 'x1'}
        templates.render_markdown(doc, {'mode': 'sections', 'columns': ['用例编号', '参考来源']}, self.out/'a.md')
        text = (self.out/'a.md').read_text()
        self.assertIn('引用：Legacy-REF?42', text)
        self.assertIn('custom_tag=x1', text)
        self.assertEqual(source_trace.normalize(doc['cases'][0]['参考来源'])[0]['custom_tag'], 'x1')
        bad = copy.deepcopy(doc)
        bad['cases'][0]['参考来源'] = {'kind': 'reference_case', 'index_url': 'recGuessedOnly'}
        template, mapping = self.builtin()
        with self.assertRaises(ValueError):
            templates.fill_xlsx(template, bad, mapping, self.out/'bad.xlsx')
        self.assertFalse((self.out/'bad.xlsx').exists())

    def test_legacy_template_compat_execution_history_and_coverage(self):
        template, mapping = self.builtin()
        doc = plain_cases()
        doc['cases'][0].update({'参考来源': {'type': '新写', 'ref': '', 'note': '依据R'}, '执行结果': '通过（2024 轮次1）'})
        mapping = dict(mapping, columns=dict(mapping['columns'], G='执行结果', H='参考来源'))
        templates.fill_xlsx(template, doc, mapping, self.out/'legacy.xlsx')
        sheet = load_workbook(self.out/'legacy.xlsx')['测试用例']
        self.assertEqual(sheet['G2'].value, '通过（2024 轮次1）')
        self.assertEqual(sheet['H2'].value, '复用方式：新写｜说明：依据R')
        self.assertEqual(compute_coverage(doc), compute_coverage(plain_cases()))
        self.assertEqual(compute_coverage(traced_cases()), compute_coverage(plain_cases()))

    def test_source_sheet_collision_and_explicit_omission(self):
        template, mapping = self.builtin()
        with self.assertRaises(ValueError):
            templates.fill_xlsx(template, traced_cases(), dict(mapping, output_sheet='来源与参考'), self.out/'c.xlsx')
        report = templates.fill_xlsx(template, traced_cases(), dict(mapping, sources_sheet=False), self.out/'o.xlsx')
        self.assertEqual(report['sources']['location'], 'omitted_by_mapping')
        report = templates.render_markdown(plain_cases(), {}, self.out/'none.md')
        self.assertEqual(report['sources']['location'], 'none_recorded')
        self.assertNotIn('来源与参考', (self.out/'none.md').read_text())


class UploadNavigationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.file = self.root/'用例.xlsx'; self.file.write_bytes(b'x')

    def tearDown(self):
        self.tmp.cleanup()

    def config(self):
        return {'config_version': '2.0', 'libraries': [{
            'name': 'formal', 'kind': 'standard', 'schema_profile': 'indexed', 'base_token': 'FormalBase01',
            'tables': {'cases': 'tblCases0001', 'review': 'tblReview001'},
            'table_base_tokens': {'review': 'MaintainBase1'},
            'field_maps': {'cases': {'用例编号': '原用例编号'}},
            'search_policy': {'default_scope': 'published', 'publication_field': '发布状态', 'published_values': ['已发布']},
            'folders': {'drafts': 'DraftFolder01', 'bodies': 'BodyFolder01', 'originals': 'OrigFolder01'}}]}

    def test_dry_run_no_request_and_draft_only_targets(self):
        with patch('lark_io.subprocess.run') as runner:
            result = write.upload_draft(self.config(), [self.file], '# x\n\nbody', title='草稿', dry_run=True,
                                        library='formal', folder_key='drafts', cases_doc=traced_cases(),
                                        backlinks=[DRAFT_PEER])
            runner.assert_not_called()
        commands = json.dumps([result['files'][0]['response']['command'], result['backlinks'][0]['command']], ensure_ascii=False)
        self.assertIn('DraftFolder01', commands)
        for formal in ('FormalBase01', 'tblCases0001', 'MaintainBase1', 'BodyFolder01', 'BodyDoc0001'):
            self.assertNotIn(formal, json.dumps(result, ensure_ascii=False).replace(INDEX_URL, '').replace(BODY_URL, ''))
        markdown = result['document']['response']['markdown']
        self.assertIn(REQ_URL, markdown); self.assertIn('TC_001_002', markdown); self.assertIn('## 草稿导航', markdown)
        self.assertEqual(result['backlinks'][0]['action'], 'planned')
        self.assertIn(write.BACKLINK_PLACEHOLDER, result['backlinks'][0]['content_preview'])

    def test_backlink_refuses_sources_formal_and_undeclared(self):
        config = self.config()
        cases_doc = traced_cases()
        formal_doc = 'https://example.invalid/docx/FormalBase01'
        cases_doc['document']['navigation']['editable_drafts'] += [BODY_URL, REQ_URL, INDEX_URL, formal_doc]
        for target in (BODY_URL, REQ_URL, INDEX_URL, formal_doc, 'https://example.invalid/docx/Undeclared01', 'PeerDraft001'):
            with patch('lark_io.subprocess.run') as runner, self.assertRaises(ValueError):
                write.upload_draft(config, [self.file], 'body', dry_run=True, library='formal', folder_key='drafts',
                                   cases_doc=cases_doc, backlinks=[target])
            runner.assert_not_called()
        with self.assertRaises(ValueError):
            write.upload_draft('DraftFolder01', [self.file], None, dry_run=True, cases_doc=cases_doc, backlinks=[DRAFT_PEER])

    def test_navigation_url_never_guessed(self):
        result = write.upload_draft('DraftFolder01', [], 'body', dry_run=True, cases_doc=plain_cases())
        markdown = result['document']['response']['markdown']
        self.assertIn('未提供草稿目录链接（不推测 URL）', markdown)
        self.assertNotIn('http', markdown)
        self.assertEqual(result['navigation']['status'], 'not_provided_not_guessed')
        result = write.upload_draft('DraftFolder01', [], 'body', dry_run=True, draft_folder_url=FOLDER_URL)
        self.assertIn('[返回草稿目录](' + FOLDER_URL + ')', result['document']['response']['markdown'])
        self.assertEqual(result['navigation']['sources_section'], 'no_cases')
        with self.assertRaises(ValueError):
            write.upload_draft('DraftFolder01', [], 'body', dry_run=True, draft_folder_url='DraftFolder01')

    def fake_remote(self, created, peer_content=''):
        calls = []
        def fake_run(args, **kwargs):
            calls.append(list(args))
            return {'data': {}}
        patches = [patch('lark_io.list_folder', side_effect=[[], [{'token': 'NewDraft0001', 'name': '草稿', 'type': 'docx'}]]),
                   patch('lark_io.create_doc', return_value=created),
                   patch('lark_io.fetch_doc', side_effect=lambda doc, **_: {'content': peer_content if doc == DRAFT_PEER else 'draft'}),
                   patch('lark_io.run_lark', side_effect=fake_run)]
        return calls, patches

    def test_real_backlink_only_appends_to_declared_draft_with_returned_url(self):
        created = {'document': {'document_id': 'NewDraft0001', 'url': 'https://example.invalid/docx/NewDraft0001'}}
        calls, patches = self.fake_remote(created)
        with contextlib.ExitStack() as stack:
            for item in patches: stack.enter_context(item)
            result = write.upload_draft('DraftFolder01', [], 'body', title='草稿', cases_doc=traced_cases(), backlinks=[DRAFT_PEER])
        self.assertEqual(result['status'], 'ok')
        self.assertEqual(result['backlinks'][0]['action'], 'appended')
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0][:6], ['docs', '+update', '--doc', DRAFT_PEER, '--command', 'append'])
        calls, patches = self.fake_remote(created, peer_content='see https://example.invalid/docx/NewDraft0001')
        with contextlib.ExitStack() as stack:
            for item in patches: stack.enter_context(item)
            result = write.upload_draft('DraftFolder01', [], 'body', title='草稿', cases_doc=traced_cases(), backlinks=[DRAFT_PEER])
        self.assertEqual(result['backlinks'][0]['action'], 'already_present')
        self.assertEqual(calls, [])

    def test_missing_draft_url_is_not_constructed(self):
        calls, patches = self.fake_remote({'document': {'document_id': 'NewDraft0001'}})
        with contextlib.ExitStack() as stack:
            for item in patches: stack.enter_context(item)
            result = write.upload_draft('DraftFolder01', [], 'body', title='草稿', cases_doc=traced_cases(), backlinks=[DRAFT_PEER])
        self.assertEqual(result['status'], 'partial')
        self.assertEqual(result['error']['code'], 'backlink_no_url')
        self.assertEqual(calls, [])

    def test_cli_upload_with_cases_dry_run(self):
        inputs = self.root/'inputs'; inputs.mkdir()
        cases_path = inputs/'cases.json'; cases_path.write_text(json.dumps(traced_cases(), ensure_ascii=False))
        md = inputs/'draft.md'; md.write_text('# 草稿\n\nbody')
        out = self.root/'report'/'upload.json'
        with patch('lark_io.subprocess.run') as runner, contextlib.redirect_stdout(io.StringIO()):
            code = write.main(['upload', '--folder', 'DraftFolder01', '--markdown', str(md), '--cases', str(cases_path),
                               '--draft-folder-url', FOLDER_URL, '--backlink', DRAFT_PEER, '--dry-run', '--out', str(out)])
            runner.assert_not_called()
        self.assertEqual(code, 0)
        report = json.loads(out.read_text())
        self.assertEqual(report['navigation'], {'draft_folder_url': FOLDER_URL, 'status': 'explicit', 'sources_section': 'appended'})


if __name__ == '__main__':
    unittest.main()
