import hashlib
import tempfile
import unittest
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from doc_extract import extract_document, scalar_fields
import doc_extract
import json
import subprocess
from datetime import time, timedelta
from unittest.mock import patch
import contextlib
import io
from lark_io import safe_output


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup); self.root = Path(self.tmp.name)

    def test_html_fake_xlsx(self):
        p = self.root/'fake.xlsx'; p.write_text('<html><p>说明</p><table><tr><td>编号</td><td>A&amp;B</td></tr></table></html>')
        before = p.read_bytes(); r = extract_document(p)
        self.assertEqual(r['format'], 'html'); self.assertEqual(r['status'], 'ok')
        self.assertTrue(any(b['location']=='表1!R1C2' and b['text']=='A&B' for b in r['blocks']))
        self.assertEqual(p.read_bytes(), before)

    def test_xlsx_merge_formula(self):
        from openpyxl import Workbook
        p = self.root/'book.xlsx'; book = Workbook(); sheet = book.active
        sheet['A1']='title'; sheet.merge_cells('A1:A3'); sheet['B2']='=1+1'; book.save(p)
        before = hashlib.sha256(p.read_bytes()).hexdigest(); r = extract_document(p)
        self.assertEqual(r['sha256'], before)
        self.assertEqual(r['blocks'][1]['merged_anchors'][0]['value'], 'title')
        self.assertEqual(r['blocks'][1]['cells'][0]['cached_value'], None)
        self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(), before)

    def test_docx_document_order(self):
        from docx import Document
        doc = Document(); doc.add_paragraph('before'); table = doc.add_table(rows=1, cols=1); table.cell(0,0).text='cell'; doc.add_paragraph('after')
        p=self.root/'doc.docx'; doc.save(p); r=extract_document(p)
        self.assertEqual([b['text'] for b in r['blocks']], ['before','cell','after'])

    def test_csv_tsv_and_text(self):
        for name, content in [('a.csv','a,"b,c"\n1,2\n'), ('a.tsv','a\tb\n'), ('a.md','# heading\ntext\n'), ('a.json','{"a":1}\n')]:
            p=self.root/name; p.write_text(content); r=extract_document(p)
            self.assertEqual(r['status'],'ok'); self.assertTrue(r['ranges'])

    def test_pdf_and_binary_not_read(self):
        for name, data in [('a.pdf',b'%PDF-1.4\n'),('a.xls',b'\xd0\xcf\x11\xe0binary'),('a.png',b'\x89PNG\x00')]:
            p=self.root/name; p.write_bytes(data); self.assertEqual(extract_document(p)['status'],'extraction_required')

    def test_output_protection_and_symlinks(self):
        source=self.root/'input'/'a.txt'; source.parent.mkdir(); source.write_text('hi')
        with self.assertRaises(ValueError):safe_output(source.parent/'out.json',[source])
        alias=self.root/'alias'; alias.symlink_to(source.parent, target_is_directory=True)
        with self.assertRaises(ValueError):safe_output(alias/'out.json',[source])
        out=self.root/'out.json'; out.write_text('keep')
        with self.assertRaises(ValueError):safe_output(out,[source])

    def test_html_spans_and_cell_paragraphs(self):
        p=self.root/'merged.xlsx'
        p.write_text('<html><table><tr><td rowspan="2" colspan="2">A</td><td>X</td></tr><tr><td><p>条件1</p><p>条件2</p></td></tr></table></html>')
        result=extract_document(p);cells={b['location']:b for b in result['blocks']}
        self.assertEqual(cells['表1!R1C1']['colspan'],2)
        self.assertEqual(cells['表1!R2C3']['text'],'条件1\n条件2')
        self.assertFalse(any(b['location']=='正文' and '条件' in b['text'] for b in result['blocks']))

    def test_time_duration_array_formula_cli(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.formula import ArrayFormula
        source=self.root/'source'/'time.xlsx';source.parent.mkdir()
        book=Workbook();sheet=book.active
        sheet['A1']=time(9,30);sheet.merge_cells('A1:B1')
        sheet['A2']=timedelta(hours=1,minutes=5);sheet['A2'].number_format='[h]:mm:ss'
        sheet['A3']=ArrayFormula(ref='A3:A4',text='=SUM(B3:B4)')
        sheet['C1']='ordinary';book.save(source)
        before=source.read_bytes();out=self.root/'out'
        run=subprocess.run([sys.executable,str(Path(doc_extract.__file__)),str(source),'--out',str(out)],capture_output=True,text=True)
        self.assertEqual(run.returncode,0,run.stderr);self.assertEqual(json.loads(run.stdout)['status'],'ok')
        result=json.loads((out/'time.xlsx.extract.json').read_text())
        cells={c['coordinate']:c for b in result['blocks'] for c in b['cells']}
        self.assertEqual((cells['A1']['value'],cells['A1']['value_type']),('09:30:00','time'))
        self.assertEqual((cells['A2']['value'],cells['A2']['value_type']),('1:05:00','timedelta'))
        self.assertEqual(cells['A3']['formula'],'=SUM(B3:B4)')
        self.assertEqual(cells['A3']['formula_type'],'ArrayFormula');self.assertEqual(cells['A3']['formula_ref'],'A3:A4')
        self.assertIn('cached_value',cells['A3']);self.assertIsNone(cells['A3']['cached_value'])
        self.assertEqual(result['blocks'][0]['merged_anchors'][0]['value_type'],'time')
        self.assertEqual(cells['C1']['value'],'ordinary');self.assertEqual(source.read_bytes(),before)
        self.assertTrue((out/'time.xlsx.extract.md').exists())

    def test_unknown_scalar_and_typed_cached_value(self):
        class Other:
            def __str__(self):return 'opaque value'
        self.assertEqual(scalar_fields(Other()),{'value':'opaque value','value_type':'Other'})
        self.assertEqual(scalar_fields(time(9), 'cached_value'),{'cached_value':'09:00:00','cached_value_type':'time'})
        json.dumps(scalar_fields(Other()))

    def test_cli_unexpected_parser_failure_is_json(self):
        source=self.root/'source'/'a.txt';source.parent.mkdir();source.write_text('hi')
        with patch.object(sys,'argv',['doc_extract',str(source),'--out',str(self.root/'out')]), patch.object(doc_extract,'extract_document',side_effect=TypeError('parser failed')),contextlib.redirect_stdout(io.StringIO()) as stdout:
            code=doc_extract.main()
        self.assertEqual(code,2);self.assertEqual(json.loads(stdout.getvalue())['status'],'failed')
        self.assertFalse((self.root/'out').exists())

    def test_html_invalid_spans_are_marked_and_located(self):
        for attr in ['rowspan','rowspan=""','colspan="2px"','colspan="0"','rowspan="-1"','rowspan="10001"']:
            with self.subTest(attr=attr):
                source=self.root/'span.html';source.write_text('<table><tr><td '+attr+'>first</td><td>second</td></tr></table>')
                result=extract_document(source);self.assertEqual(result['status'],'ok')
                self.assertEqual([b['location'] for b in result['blocks']],['表1!R1C1','表1!R1C2'])
                self.assertEqual(result['blocks'][0]['span_adjustments'][0]['used'],1)
                self.assertEqual(result['blocks'][1]['text'],'second');json.dumps(result)
